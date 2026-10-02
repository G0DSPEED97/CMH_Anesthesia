"""Build the manager-facing project report from docs/PROJECT_REPORT.md."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PROJECT_REPORT.md"
OUTPUT = ROOT / "docs" / "reports" / "CMH_Anaesthesia_Project_Experience_and_Pilot_Readiness_Report.docx"

GREEN = "164A3D"
PALE = "F2F6F3"
GRID = "D9D9D9"
TEXT = RGBColor(22, 33, 29)
MUTED = RGBColor(92, 106, 100)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=130, bottom=120, end=130) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color: str = GRID, size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        else:
            paragraph.add_run(part)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(22)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 20, 16, 8),
        ("Heading 2", 15, 13, 6),
        ("Heading 3", 12, 10, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Figure Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False

    header = section.header.paragraphs[0]
    header.text = "CMH Anaesthesia Management System"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths_by_columns = {
        3: [1.35, 2.65, 2.8],
        4: [0.85, 1.85, 2.2, 1.9],
    }
    widths = widths_by_columns.get(len(rows[0]), [6.75 / len(rows[0])] * len(rows[0]))
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if row_index == 0:
                set_cell_shading(cell, GREEN)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.6 if len(rows[0]) == 4 else 9.2)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        parts = [item.strip() for item in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", item) for item in parts):
            rows.append(parts)
        i += 1
    return rows, i


def build() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)

    i = 0
    on_cover = True
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line == "<!-- pagebreak -->":
            doc.add_page_break()
            on_cover = False
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        image_match = re.fullmatch(r"!\[([^]]+)]\(([^)]+)\)", line)
        if image_match:
            caption, relative = image_match.groups()
            image_path = ROOT / "docs" / relative
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(image_path), width=Inches(6.72))
            cap = doc.add_paragraph(style="Figure Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.add_run(caption)
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            value = heading.group(2)
            paragraph = doc.add_paragraph(style="Title" if level == 1 else f"Heading {level - 1}")
            paragraph.add_run(value)
            if level == 1:
                paragraph.paragraph_format.space_before = Pt(92)
            i += 1
            continue
        if re.match(r"^-\s+", line):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.27)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline(paragraph, "•  " + re.sub(r"^-\s+", "", line))
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.first_line_indent = Inches(-0.24)
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline(paragraph, f"{numbered.group(1)}.  {numbered.group(2)}")
            i += 1
            continue

        if on_cover and line.startswith("**"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(5)
            add_inline(paragraph, line)
            i += 1
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if not candidate or candidate.startswith(("#", "|", "- ", "![", "<!--")) or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            i += 1
        paragraph = doc.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines).replace("  ", " "))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "CMH Anaesthesia Management System Project Experience and Pilot Readiness Report"
    doc.core_properties.subject = "Project history, delivered workflow, database controls, and pilot decisions"
    doc.core_properties.author = "CMH Anaesthesia project team"
    doc.core_properties.keywords = "CMH, anaesthesia, project report, PostgreSQL, Alembic, pilot"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
